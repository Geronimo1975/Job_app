"""
Excel Tracker — Management centralizat al aplicațiilor în Excel.

Gestionează Job_Application_Tracker.xlsx:
  - Adăugare joburi noi
  - Actualizare status
  - Căutare / filtrare
  - Backup automat
"""

import shutil
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# Coloanele tracker-ului (ordine fixă)
COLUMNS = {
    "A": ("Nr.", 6),
    "B": ("Firma", 22),
    "C": ("Position", 28),
    "D": ("Standort", 16),
    "E": ("URL", 30),
    "F": ("Quelle", 12),
    "G": ("Kontakt Email", 28),
    "H": ("Kontakt Person", 18),
    "I": ("Sprache", 8),
    "J": ("Gefunden am", 12),
    "K": ("Status", 16),
    "L": ("Anschreiben", 20),
    "M": ("Match Score", 10),
    "N": ("Keywords Match", 24),
    "O": ("Email Inhalt", 30),
    "P": ("Gesendet am", 14),
    "Q": ("Email Status", 14),
    "R": ("Anschreiben PDF", 30),
    "S": ("CV PDF", 30),
    "T": ("Zeugnisse", 24),
    "U": ("Zertifikate", 24),
    "V": ("Komplett PDF", 28),
    "W": ("Sende Befehl", 14),
    "X": ("Notizen", 30),
}

# Status-uri posibile
STATUS_COLORS = {
    "Neu": "E3F2FD",          # light blue
    "Analysiert": "FFF3E0",   # light orange
    "Anschreiben": "F3E5F5",  # light purple
    "Bereit": "E8F5E9",      # light green
    "SENDEN": "FFEB3B",      # yellow
    "GESENDET": "4CAF50",    # green
    "Antwort": "2196F3",     # blue
    "Interview": "FF9800",   # orange
    "Absage": "F44336",      # red
    "FEHLER": "F44336",      # red
}


class ExcelTracker:
    """Manager pentru Excel Job Tracker."""

    def __init__(self, config: dict, project_root: Path):
        self.config = config
        self.project_root = project_root
        tracker_cfg = config.get("tracker", {})
        self.tracker_path = project_root / tracker_cfg.get(
            "file", "data/Job_Application_Tracker.xlsx"
        )
        self.auto_backup = tracker_cfg.get("auto_backup", True)
        self.backup_dir = project_root / tracker_cfg.get("backup_dir", "data/backups")

        if not HAS_OPENPYXL:
            raise ImportError("openpyxl nu este instalat. Rulează: pip install openpyxl")

    def _ensure_tracker(self) -> openpyxl.Workbook:
        """Creează tracker-ul dacă nu există, sau îl deschide."""
        if self.tracker_path.exists():
            return openpyxl.load_workbook(self.tracker_path)

        # Creează tracker nou
        self.tracker_path.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Bewerbungen"

        # Header row
        header_fill = PatternFill("solid", fgColor="1565C0")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col_letter, (name, width) in COLUMNS.items():
            cell = ws[f"{col_letter}1"]
            cell.value = name
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[col_letter].width = width

        # Freeze header
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:X1"

        wb.save(self.tracker_path)
        return wb

    def backup(self) -> Path | None:
        """Creează un backup al tracker-ului."""
        if not self.tracker_path.exists():
            return None

        self.backup_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = self.backup_dir / f"Tracker_Backup_{timestamp}.xlsx"
        shutil.copy2(self.tracker_path, backup_path)
        return backup_path

    def add_job(self, job_data: dict) -> int:
        """
        Adaugă un job nou în tracker.

        Args:
            job_data: Dict cu datele jobului (company, position, url, etc.)

        Returns:
            Numărul rândului adăugat.
        """
        wb = self._ensure_tracker()
        ws = wb.active

        if self.auto_backup:
            self.backup()

        # Determină următorul rând
        next_row = ws.max_row + 1
        nr = next_row - 1  # Nr. = row - header

        # Verifică duplicat (după URL sau companie+poziție)
        url = job_data.get("url", "")
        company = job_data.get("company", "")
        position = job_data.get("title", job_data.get("position", ""))

        for row in range(2, next_row):
            existing_url = ws.cell(row=row, column=5).value or ""
            existing_company = ws.cell(row=row, column=2).value or ""
            existing_position = ws.cell(row=row, column=3).value or ""
            if (url and existing_url == url) or (
                existing_company == company and existing_position == position
            ):
                return row  # Deja există

        # Scrie datele
        ws.cell(row=next_row, column=1, value=nr)                                    # Nr.
        ws.cell(row=next_row, column=2, value=company)                                # Firma
        ws.cell(row=next_row, column=3, value=position)                               # Position
        ws.cell(row=next_row, column=4, value=job_data.get("location", ""))           # Standort
        ws.cell(row=next_row, column=5, value=url)                                    # URL
        ws.cell(row=next_row, column=6, value=job_data.get("source", ""))             # Quelle
        ws.cell(row=next_row, column=7, value=job_data.get("email_hr", ""))           # Kontakt Email
        ws.cell(row=next_row, column=8, value=job_data.get("contact_person", ""))     # Kontakt Person
        ws.cell(row=next_row, column=9, value=job_data.get("language", "de"))         # Sprache
        ws.cell(row=next_row, column=10, value=datetime.now().strftime("%Y-%m-%d"))   # Gefunden am
        ws.cell(row=next_row, column=11, value="Neu")                                 # Status

        # Stil
        self._apply_status_style(ws, next_row, "Neu")

        wb.save(self.tracker_path)
        return next_row

    def update_status(self, row: int, status: str, **extra_fields) -> None:
        """Actualizează statusul unui job în tracker."""
        wb = self._ensure_tracker()
        ws = wb.active

        ws.cell(row=row, column=11, value=status)
        self._apply_status_style(ws, row, status)

        # Câmpuri extra (column_index: value)
        field_map = {
            "email_status": 17,
            "sende_befehl": 23,
            "anschreiben_pdf": 18,
            "cv_pdf": 19,
            "gesendet_am": 16,
            "notizen": 24,
            "match_score": 13,
            "keywords_match": 14,
            "email_inhalt": 15,
            "komplett_pdf": 22,
        }

        for key, val in extra_fields.items():
            if key in field_map:
                ws.cell(row=row, column=field_map[key], value=val)

        wb.save(self.tracker_path)

    def get_jobs_by_status(self, status: str) -> list[dict]:
        """Returnează joburi cu un anumit status."""
        wb = self._ensure_tracker()
        ws = wb.active
        jobs = []

        for row in range(2, ws.max_row + 1):
            cell_status = ws.cell(row=row, column=11).value or ""
            if cell_status.upper() == status.upper():
                jobs.append(self._row_to_dict(ws, row))

        return jobs

    def get_jobs_to_send(self) -> list[dict]:
        """Returnează joburi cu Sende Befehl = SENDEN."""
        wb = self._ensure_tracker()
        ws = wb.active
        jobs = []

        for row in range(2, ws.max_row + 1):
            sende = ws.cell(row=row, column=23).value or ""
            email = ws.cell(row=row, column=7).value or ""
            if sende.upper() == "SENDEN" and email:
                jobs.append(self._row_to_dict(ws, row))

        return jobs

    def get_all_jobs(self) -> list[dict]:
        """Returnează toate joburile din tracker."""
        wb = self._ensure_tracker()
        ws = wb.active
        return [self._row_to_dict(ws, row) for row in range(2, ws.max_row + 1)]

    def _row_to_dict(self, ws, row: int) -> dict:
        """Convertește un rând din tracker în dict."""
        col_keys = list(COLUMNS.keys())
        col_names = [v[0] for v in COLUMNS.values()]
        data = {"_row": row}
        for i, (letter, (name, _)) in enumerate(COLUMNS.items()):
            val = ws.cell(row=row, column=i + 1).value
            data[name] = str(val).strip() if val else ""
        return data

    def _apply_status_style(self, ws, row: int, status: str) -> None:
        """Aplică culoarea corespunzătoare statusului."""
        color = STATUS_COLORS.get(status, "FFFFFF")
        cell = ws.cell(row=row, column=11)
        cell.fill = PatternFill("solid", fgColor=color)

        if status in ("GESENDET", "FEHLER", "Absage"):
            cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        else:
            cell.font = Font(name="Arial", size=9)
